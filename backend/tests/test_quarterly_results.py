from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.export_xlsx import quarterly_results_workbook, workbook_bytes
from app.services.quarterly_results import filter_results_clients


def test_filter_results_clients():
    rows = [
        {"counterparty": "ИП Almaz-A", "work_type_label": "Удержание", "manager_name": "Иванов"},
        {"counterparty": "Гранат", "work_type_label": "Рост", "manager_name": "Петров"},
    ]
    assert [c["counterparty"] for c in filter_results_clients(rows, query="almaz")] == ["ИП Almaz-A"]
    assert [c["counterparty"] for c in filter_results_clients(rows, work_type="рост")] == ["Гранат"]
    assert [c["counterparty"] for c in filter_results_clients(rows, manager="петр")] == ["Гранат"]
    assert len(filter_results_clients(rows)) == 2


def test_quarterly_results_workbook_columns():
    report = {
        "labels": {
            "plan": "План отгрузок на 3 квартал",
            "shipment_fact": "Факт отгрузок 3 квартал",
            "shipment_percent": "% выполнения",
            "shipment_prev": "Факт отгрузок 2 квартал",
            "shipment_prev2": "Факт отгрузок 1 квартал",
            "shipment_dynamics": "Динамика отгрузок",
            "sales": "Продажи 3 кв.",
            "sales_prev": "Продажи 2 кв.",
            "sales_prev2": "Продажи 1 кв.",
            "sales_dynamics": "Динамика продаж",
        },
        "clients": [
            {
                "counterparty": "ИП Garant.S",
                "manager_name": "Иванов",
                "work_type_label": "Удержание",
                "work_type_percent": 0,
                "plan": 0,
                "shipment_fact": 40,
                "shipment_percent": 0,
                "shipment_prev_quarter": 50,
                "shipment_prev2_quarter": 30,
                "shipment_dynamics_percent": 80,
                "sales_total": 34,
                "sales_prev_quarter": 80,
                "sales_prev2_quarter": 70,
                "dynamics_percent": 42.5,
                "comment": "Участвует",
            }
        ],
    }
    data = workbook_bytes(quarterly_results_workbook(report))
    assert data[:2] == b"PK"
    ws = load_workbook(BytesIO(data)).active
    assert ws["A1"].value == "Головной контрагент"
    assert ws["E1"].value == "План отгрузок на 3 квартал"
    assert ws["O1"].value == "Комментарий"
    assert ws["A2"].value == "ИП Garant.S"
    assert ws["E2"].value == 0
    assert ws["N2"].value == 42.5
    assert ws["O2"].value == "Участвует"


def test_zero_fact_placeholder():
    from uuid import uuid4

    from app.services.quarterly_results import _zero_fact

    cp = SimpleNamespace(id=uuid4(), name="Клиент")
    fact = _zero_fact(cp, 2026, 3)
    assert fact.fact_amount == 0
    assert fact.counterparty == "Клиент"
