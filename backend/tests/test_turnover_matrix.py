from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4

from app.domain.articles import index_nomenclature
from app.domain.turnover_matrix import (
    SKU_ROW_LIMIT,
    article_stock_begin,
    assemble_turnover_rows,
    filter_empty_turnover_rows,
    header_stock_begin,
    is_empty_turnover_row,
    month_cell,
)
from app.services.turnover_matrix import _month_iter


def test_month_iter_span():
    months = _month_iter(2023, 11, 2024, 2)
    assert months == [(2023, 11), (2023, 12), (2024, 1), (2024, 2)]


def test_turnover_matrix_workbook_includes_percent():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.export_xlsx import turnover_matrix_workbook, workbook_bytes

    report = {
        "view": "lts",
        "months": ["2026-07"],
        "rows": [
            {
                "counterparty": "ИП Тест",
                "months": {
                    "2026-07": {
                        "stock_begin": 10,
                        "stock_end": 10,
                        "sales": 10,
                        "turnover_percent": 100,
                    }
                },
            }
        ],
    }
    wb = load_workbook(BytesIO(workbook_bytes(turnover_matrix_workbook(report))))
    sheet = wb.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "2026-07 Об-ть %" in headers
    assert sheet.cell(2, headers.index("2026-07 Об-ть %") + 1).value == 100


def test_header_begin_is_zero_without_history():
    assert header_stock_begin(Decimal(0)) == Decimal(0)


def test_article_begin_falls_back_to_end():
    assert article_stock_begin(Decimal(0), False, Decimal(9)) == Decimal(9)
    assert article_stock_begin(Decimal(4), True, Decimal(9)) == Decimal(4)


def test_month_cell_turnover():
    cell = month_cell(Decimal(10), Decimal(10), Decimal(10))
    assert cell["sales"] == 10
    assert cell["stock_avg"] == 10
    assert cell["turnover_percent"] == 100


def test_empty_row_detects_zeros_including_movements():
    zero = {
        "months": {
            "2026-07": {
                "stock_begin": 0,
                "stock_end": 0,
                "sales": 0,
                "realization": 0,
                "return_qty": 0,
            }
        }
    }
    live = {"months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 2}}}
    assert is_empty_turnover_row(zero) is True
    assert is_empty_turnover_row(live) is False
    assert is_empty_turnover_row({"months": {}}) is True


def test_filter_empty_keeps_parent_if_child_has_qty():
    parent = {
        "row_type": "counterparty",
        "counterparty": "A",
        "months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 0}},
    }
    empty_child = {
        "row_type": "dimension",
        "dimension": "Архив",
        "months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 0}},
    }
    live_child = {
        "row_type": "dimension",
        "dimension": "Актив",
        "months": {"2026-07": {"stock_begin": 1, "stock_end": 0, "sales": 0}},
    }
    rows = filter_empty_turnover_rows([parent, empty_child, live_child])
    assert [r.get("dimension") or r.get("counterparty") for r in rows] == ["A", "Актив"]


def test_filter_empty_drops_all_zero_group():
    parent = {
        "row_type": "counterparty",
        "counterparty": "A",
        "months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 0}},
    }
    child = {
        "row_type": "sku",
        "article": "X",
        "months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 0}},
    }
    assert filter_empty_turnover_rows([parent, child]) == []


def test_filter_empty_keeps_parent_with_qty_even_if_children_zero():
    parent = {
        "row_type": "counterparty",
        "counterparty": "A",
        "months": {"2026-07": {"stock_begin": 0, "stock_end": 2, "sales": 0}},
    }
    child = {
        "row_type": "dimension",
        "dimension": "Архив",
        "months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 0}},
    }
    rows = filter_empty_turnover_rows([parent, child])
    assert len(rows) == 1
    assert rows[0]["counterparty"] == "A"
    zero = {"counterparty": "Z", "months": {"2026-07": {"stock_begin": 0, "stock_end": 0, "sales": 0}}}
    live = {"counterparty": "L", "months": {"2026-07": {"stock_begin": 0, "stock_end": 3, "sales": 0}}}
    assert [r["counterparty"] for r in filter_empty_turnover_rows([zero, live])] == ["L"]


def _bounds():
    return [("2026-07", date(2026, 7, 1), date(2026, 7, 31))]


def test_assemble_lts_groups_and_header_begin_zero():
    cp_id = uuid4()
    cp = SimpleNamespace(id=cp_id, name="ИП Тест", work_type="hold", work_type_percent=0)
    nom = SimpleNamespace(
        id=uuid4(),
        article="A1",
        barcode=None,
        name="Кольцо",
        lts="Актив",
        wear_type="Кольцо",
        metal_color="Красное",
        lts_date=None,
    )
    sales = [
        SimpleNamespace(head_counterparty_id=cp_id, article="A1", quantity=Decimal(4), period_year=2026, period_month=7)
    ]
    stocks = [
        SimpleNamespace(head_counterparty_id=cp_id, article="A1", quantity=Decimal(10), stock_date=date(2026, 7, 31))
    ]
    rows = assemble_turnover_rows(
        view="lts",
        month_bounds=_bounds(),
        counterparties=[cp],
        sales=sales,
        stocks=stocks,
        noms=index_nomenclature([nom]),
    )
    assert rows[0]["row_type"] == "counterparty"
    assert rows[0]["months"]["2026-07"]["sales"] == 4
    assert rows[0]["months"]["2026-07"]["stock_begin"] == 0
    assert rows[0]["months"]["2026-07"]["stock_end"] == 10
    child = rows[1]
    assert child["dimension"] == "Актив"
    assert child["months"]["2026-07"]["stock_begin"] == 10
    assert child["months"]["2026-07"]["sales"] == 4


def test_assemble_sku_uses_movements_and_limit():
    cp_id = uuid4()
    nom_id = uuid4()
    cp = SimpleNamespace(id=cp_id, name="ИП Тест", work_type="рост", work_type_percent=Decimal(10))
    nom = SimpleNamespace(
        id=nom_id,
        article="SKU-1",
        barcode=None,
        name="Серьги",
        lts="Актив",
        wear_type="Серьги",
        metal_color="Белое",
        lts_date=date(2025, 1, 15),
    )
    sales = [
        SimpleNamespace(
            head_counterparty_id=cp_id, article="SKU-1", quantity=Decimal(2), period_year=2026, period_month=7
        )
    ]
    stocks = [
        SimpleNamespace(head_counterparty_id=cp_id, article="SKU-1", quantity=Decimal(5), stock_date=date(2026, 6, 30)),
        SimpleNamespace(head_counterparty_id=cp_id, article="SKU-1", quantity=Decimal(3), stock_date=date(2026, 7, 31)),
    ]
    movements = {(cp_id, nom_id, 2026, 7): (Decimal(8), Decimal(1))}
    rows = assemble_turnover_rows(
        view="main",
        month_bounds=_bounds(),
        counterparties=[cp],
        sales=sales,
        stocks=stocks,
        noms=index_nomenclature([nom]),
        movements=movements,
    )
    assert rows[0]["row_type"] == "counterparty"
    sku = rows[1]
    assert sku["article"] == "SKU-1"
    assert sku["lts_date"] == "2025-01-15"
    cell = sku["months"]["2026-07"]
    assert cell["stock_begin"] == 5
    assert cell["stock_end"] == 3
    assert cell["sales"] == 2
    assert cell["realization"] == 8
    assert cell["return_qty"] == 1


def test_sku_row_limit_constant():
    assert SKU_ROW_LIMIT == 500


def test_sku_row_limit_truncates():
    cp_id = uuid4()
    cp = SimpleNamespace(id=cp_id, name="ИП Тест", work_type=None, work_type_percent=None)
    sales = [
        SimpleNamespace(
            head_counterparty_id=cp_id,
            article=f"A{i:04d}",
            quantity=Decimal(1),
            period_year=2026,
            period_month=7,
        )
        for i in range(SKU_ROW_LIMIT + 10)
    ]
    rows = assemble_turnover_rows(
        view="main",
        month_bounds=_bounds(),
        counterparties=[cp],
        sales=sales,
        stocks=[],
        noms={},
    )
    assert len([r for r in rows if r.get("row_type") == "sku"]) == SKU_ROW_LIMIT


def test_empty_helpers_and_counterparty_view():
    from app.domain.turnover_matrix import empty_month_cell, is_empty_month_cell

    assert is_empty_month_cell(None) is True
    assert empty_month_cell(main=True)["realization"] == 0
    assert month_cell(None, None, None)["sales"] == 0  # type: ignore[arg-type]
    cp_id = uuid4()
    cp = SimpleNamespace(id=cp_id, name="ИП Тест", work_type=None, work_type_percent=None)
    rows = assemble_turnover_rows(
        view="counterparty",
        month_bounds=_bounds(),
        counterparties=[cp],
        sales=[
            SimpleNamespace(
                head_counterparty_id=cp_id, article="A1", quantity=Decimal(3), period_year=2026, period_month=7
            )
        ],
        stocks=[],
        noms={},
    )
    assert rows[0]["counterparty"] == "ИП Тест"
    assert rows[0]["months"]["2026-07"]["sales"] == 3
    assert "row_type" not in rows[0]


def test_dimension_fills_missing_month_and_uses_begin_snapshot():
    cp_id = uuid4()
    cp = SimpleNamespace(id=cp_id, name="ИП Тест", work_type=None, work_type_percent=0)
    nom = SimpleNamespace(
        id=uuid4(),
        article="A1",
        barcode=None,
        name="Кольцо",
        lts="Актив",
        wear_type="Кольцо",
        metal_color="Красное",
        lts_date=None,
    )
    rows = assemble_turnover_rows(
        view="lts",
        month_bounds=[
            ("2026-07", date(2026, 7, 1), date(2026, 7, 31)),
            ("2026-08", date(2026, 8, 1), date(2026, 8, 31)),
        ],
        counterparties=[cp],
        sales=[
            SimpleNamespace(
                head_counterparty_id=cp_id, article="A1", quantity=Decimal(1), period_year=2026, period_month=7
            )
        ],
        stocks=[
            SimpleNamespace(head_counterparty_id=cp_id, article="A1", quantity=Decimal(6), stock_date=date(2026, 6, 30)),
            SimpleNamespace(head_counterparty_id=cp_id, article="A1", quantity=Decimal(4), stock_date=date(2026, 7, 31)),
        ],
        noms=index_nomenclature([nom]),
    )
    child = rows[1]
    assert child["months"]["2026-07"]["stock_begin"] == 6
    assert child["months"]["2026-08"]["sales"] == 0
    assert child["months"]["2026-08"]["stock_begin"] == 10


def test_unknown_article_goes_to_dash_dimension():
    cp_id = uuid4()
    cp = SimpleNamespace(id=cp_id, name="ИП Тест", work_type=None, work_type_percent=0)
    sales = [
        SimpleNamespace(
            head_counterparty_id=cp_id, article="NOPE", quantity=Decimal(1), period_year=2026, period_month=7
        )
    ]
    rows = assemble_turnover_rows(
        view="wear_type",
        month_bounds=_bounds(),
        counterparties=[cp],
        sales=sales,
        stocks=[],
        noms={},
    )
    assert rows[1]["dimension"] == "—"


def test_build_matrix_empty_period_skips_db():
    from app.services.turnover_matrix import build_turnover_matrix

    class Boom:
        def scalars(self, *_a, **_k):
            raise AssertionError("must not query")

    report = build_turnover_matrix(
        Boom(),
        view="lts",
        year_from=2026,
        month_from=9,
        year_to=2026,
        month_to=7,
    )
    assert report["rows"] == []
    assert report["months"] == []


def test_build_matrix_skips_empty_scope():
    from app.services.turnover_matrix import build_turnover_matrix

    class Boom:
        def scalars(self, *_a, **_k):
            raise AssertionError("must not query")

    report = build_turnover_matrix(
        Boom(),
        view="lts",
        year_from=2026,
        month_from=7,
        year_to=2026,
        month_to=9,
        allowed_ids=set(),
    )
    assert report["rows"] == []
    assert report["months"] == ["2026-07", "2026-08", "2026-09"]


def test_load_helpers_skip_empty_ids():
    from app.services.turnover_matrix import _load_movements, _load_sales, _load_stocks

    class Boom:
        def scalars(self, *_a, **_k):
            raise AssertionError("must not query")

    assert _load_sales(Boom(), [], [(2026, 7)], view="lts", year_from=2026, year_to=2026) == []
    assert _load_stocks(Boom(), []) == []
    assert _load_movements(Boom(), [], set(), date(2026, 7, 1), date(2026, 7, 31)) == {}
